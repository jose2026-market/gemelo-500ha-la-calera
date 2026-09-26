# -*- coding: utf-8 -*-
"""PPT de 4 diapositivas para productores · Sociedad Rural de San Luis."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
from pathlib import Path

ROOT = Path(__file__).resolve().parent
FOTO = ROOT / "assets" / "hereford-vaca-ternero-belgrano.png"
LOGO = ROOT / "assets" / "logo-sociedad-rural-san-luis.png"
OUT = ROOT / "Gemelo-Digital-Ganadero-productores.pptx"
URL = "https://jose2026-market.github.io/gemelo-500ha-la-calera/"

VERDE = RGBColor(0x0D, 0x5C, 0x2C)
VERDE_OSC = RGBColor(0x08, 0x3A, 0x1C)
ORO = RGBColor(0xC4, 0xA3, 0x5A)
CREMA = RGBColor(0xF7, 0xF1, 0xE4)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
TIERRA = RGBColor(0x2D, 0x20, 0x16)
CREMA2 = RGBColor(0xE8, 0xF0, 0xE4)

W, H = Inches(13.333), Inches(7.5)


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def box(slide, l, t, w, h, color, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    fill(s, color)
    s.adjustments[0] = 0.08
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    else:
        s.line.fill.background()
    return s


def txt(shape, text, size, color, bold=False, align=PP_ALIGN.LEFT, font="Calibri"):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run() if p.runs else p.add_run()
    # clear default
    p.clear()
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return tf


def para(tf, text, size, color, bold=False, align=PP_ALIGN.LEFT, space=6):
    p = tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space)
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return p


def set_bg(slide, color):
    bg = slide.background
    fill_el = bg._element
    # solid fill via spPr
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def appear_on_click(slide, shape, order):
    """Entrance: appear when the presenter clicks (order 1, 2, 3...)."""
    spid = str(shape._element.get("id"))
    ns = {
        "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }
    sld = slide._element
    timing = sld.find(qn("p:timing"))
    if timing is None:
        timing = etree.SubElement(sld, qn("p:timing"))
        tnLst = etree.SubElement(timing, qn("p:tnLst"))
        par = etree.SubElement(tnLst, qn("p:par"))
        cTn = etree.SubElement(par, qn("p:cTn"))
        cTn.set("id", "1")
        cTn.set("dur", "indefinite")
        cTn.set("restart", "never")
        cTn.set("nodeType", "tmRoot")
        child = etree.SubElement(cTn, qn("p:childTnLst"))
        seq = etree.SubElement(child, qn("p:seq"))
        seq.set("concurrent", "1")
        seq.set("nextAc", "seek")
        cTn2 = etree.SubElement(seq, qn("p:cTn"))
        cTn2.set("id", "2")
        cTn2.set("dur", "indefinite")
        cTn2.set("nodeType", "mainSeq")
        etree.SubElement(cTn2, qn("p:childTnLst"))
        prevCondLst = etree.SubElement(seq, qn("p:prevCondLst"))
        cond = etree.SubElement(prevCondLst, qn("p:cond"))
        cond.set("evt", "onPrev")
        cond.set("delay", "0")
        tgtEl = etree.SubElement(cond, qn("p:tgtEl"))
        etree.SubElement(tgtEl, qn("p:sldTgt"))
        nextCondLst = etree.SubElement(seq, qn("p:nextCondLst"))
        cond2 = etree.SubElement(nextCondLst, qn("p:cond"))
        cond2.set("evt", "onNext")
        cond2.set("delay", "0")
        tgtEl2 = etree.SubElement(cond2, qn("p:tgtEl"))
        etree.SubElement(tgtEl2, qn("p:sldTgt"))

    main = sld.find(".//" + qn("p:cTn") + "[@nodeType='mainSeq']/" + qn("p:childTnLst"))
    if main is None:
        return
    uid = 20 + order * 10
    xml = f"""
    <p:par xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
      <p:cTn id="{uid}" fill="hold">
        <p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>
        <p:childTnLst>
          <p:par>
            <p:cTn id="{uid+1}" fill="hold">
              <p:stCondLst><p:cond delay="0"/></p:stCondLst>
              <p:childTnLst>
                <p:par>
                  <p:cTn id="{uid+2}" presetID="1" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="clickEffect">
                    <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                    <p:childTnLst>
                      <p:set>
                        <p:cBhvr>
                          <p:cTn id="{uid+3}" dur="1" fill="hold">
                            <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                          </p:cTn>
                          <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
                          <p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>
                        </p:cBhvr>
                        <p:to><p:strVal val="visible"/></p:to>
                      </p:set>
                    </p:childTnLst>
                  </p:cTn>
                </p:par>
              </p:childTnLst>
            </p:cTn>
          </p:par>
        </p:childTnLst>
      </p:cTn>
    </p:par>
    """
    main.append(etree.fromstring(xml))


def nav_btn(slide, label, left, top, w=Inches(1.35), h=Inches(0.38)):
    s = box(slide, left, top, w, h, ORO)
    s.text_frame.word_wrap = False
    p = s.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.clear()
    r = p.add_run()
    r.text = label
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = VERDE_OSC
    r.font.name = "Calibri"
    s.text_frame.paragraphs[0].space_before = Pt(0)
    try:
        s.text_frame.auto_size = None
    except Exception:
        pass
    return s


def header(slide, kicker):
    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), Inches(12.35), Inches(0.22), Inches(0.72), Inches(0.72))
    s = box(slide, Inches(0.45), Inches(0.28), Inches(7.2), Inches(0.42), VERDE_OSC)
    fill(s, VERDE_OSC)
    tf = s.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.clear()
    r = p.add_run()
    r.text = kicker
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = ORO
    r.font.name = "Calibri"


def title(slide, text, top=Inches(0.85)):
    s = slide.shapes.add_textbox(Inches(0.5), top, Inches(8.2), Inches(1.35))
    tf = s.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.clear()
    r = p.add_run()
    r.text = text
    r.font.size = Pt(40)
    r.font.bold = True
    r.font.color.rgb = BLANCO
    r.font.name = "Georgia"
    return s


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    # ——— 1 ———
    s1 = prs.slides.add_slide(blank)
    set_bg(s1, VERDE_OSC)
    header(s1, "SOCIEDAD RURAL DE SAN LUIS  ·  SERVICIO AL PRODUCTOR  ·  DPTO. BELGRANO")
    title(s1, "Su campo.\nEn el mapa.")
    baj = s1.shapes.add_textbox(Inches(0.5), Inches(2.35), Inches(6.6), Inches(1.35))
    tf = baj.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.clear()
    r = p.add_run()
    r.text = "No es un papel. Es su lote en el satélite, con suelo INTA, relieve y aguadas, para decidir."
    r.font.size = Pt(18)
    r.font.color.rgb = CREMA
    r.font.name = "Georgia"

    bullets = [
        "Se abre un campo de ejemplo en Belgrano. Después, el del productor.",
        "Belgrano: 53.117 bovinos · 94,6 % de los campos son pecuarios.",
        "Hoy la zona rinde ~6 kg carne/ha. Con manejo, se puede duplicar.",
        "Sin desmonte. Sin recortar el agua de bebida. No es soja.",
    ]
    bl = s1.shapes.add_textbox(Inches(0.5), Inches(3.85), Inches(6.7), Inches(2.4))
    tf = bl.text_frame
    tf.word_wrap = True
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if i == 0:
            p.clear()
        p.space_before = Pt(6)
        r = p.add_run()
        r.text = "▸  " + line
        r.font.size = Pt(15)
        r.font.color.rgb = CREMA
        r.font.name = "Calibri"

    if FOTO.exists():
        s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.45), Inches(1.15), Inches(5.4), Inches(5.05))
        pic_bg = s1.shapes[-1]
        fill(pic_bg, ORO)
        pic_bg.adjustments[0] = 0.04
        s1.shapes.add_picture(str(FOTO), Inches(7.58), Inches(1.28), Inches(5.14), Inches(4.45))
        cap = s1.shapes.add_textbox(Inches(7.55), Inches(5.78), Inches(5.2), Inches(0.45))
        tf = cap.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.clear()
        r = p.add_run()
        r.text = "Hereford con ternero  ·  cría típica del oeste de San Luis"
        r.font.size = Pt(11)
        r.font.italic = True
        r.font.color.rgb = ORO
        r.font.name = "Calibri"

    n1 = nav_btn(s1, "Siguiente  →", Inches(11.5), Inches(6.92))

    # ——— 2 ———
    s2 = prs.slides.add_slide(blank)
    set_bg(s2, VERDE_OSC)
    header(s2, "INTA REGIÓN III  ·  INICIAL  ·  MEJORANDO  ·  IDEAL  ·  AÑO COMPLETO")
    title(s2, "Tres situaciones.", top=Inches(0.82))
    sub = s2.shapes.add_textbox(Inches(0.5), Inches(1.55), Inches(12.2), Inches(0.4))
    tf = sub.text_frame
    p = tf.paragraphs[0]
    p.clear()
    r = p.add_run()
    r.text = "El invierno limita cuántos animales se mantienen. Números INTA de zona; en su campo se lee ese lote."
    r.font.size = Pt(15)
    r.font.color.rgb = CREMA
    r.font.name = "Calibri"

    cards = [
        ("INICIAL", "12,6 ha/EV", "Destete ~62 %\n~5,9 kg carne/ha\nCómo está hoy el lote."),
        ("MEJORANDO", "10 ha/EV", "Destete ~85 %\n~12 kg/ha\nRotar, dividir, aguada cerca.\nSembrar 1–2 ha. Sin desmonte."),
        ("IDEAL", "Techo INTA", "Destete ~90 % en bajo\nEn loma no pasa de Mejorando.\nBuffel solo si OTBN lo permite."),
    ]
    xs = [0.5, 4.7, 8.9]
    card_shapes = []
    for x, (tit, num, body) in zip(xs, cards):
        c = box(s2, Inches(x), Inches(2.15), Inches(3.9), Inches(3.85), CREMA, ORO)
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.18)
        tf.margin_right = Inches(0.14)
        tf.margin_top = Inches(0.16)
        p = tf.paragraphs[0]
        p.clear()
        r = p.add_run()
        r.text = tit
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = VERDE
        r.font.name = "Calibri"
        p2 = tf.add_paragraph()
        p2.space_before = Pt(10)
        r2 = p2.add_run()
        r2.text = num
        r2.font.size = Pt(26)
        r2.font.bold = True
        r2.font.color.rgb = VERDE_OSC
        r2.font.name = "Georgia"
        p3 = tf.add_paragraph()
        p3.space_before = Pt(10)
        r3 = p3.add_run()
        r3.text = body
        r3.font.size = Pt(14)
        r3.font.color.rgb = TIERRA
        r3.font.name = "Calibri"
        card_shapes.append(c)

    regla = s2.shapes.add_textbox(Inches(0.5), Inches(6.12), Inches(10.2), Inches(0.55))
    tf = regla.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.clear()
    r = p.add_run()
    r.text = "No se recorta el agua de bebida.  ·  No se desmonta.  ·  No es soja.  ·  Círculo continuo 400 m = cría; de puntos 800 m = adulta."
    r.font.size = Pt(13)
    r.font.color.rgb = ORO
    r.font.name = "Calibri"

    btn_back2 = nav_btn(s2, "←  Atrás", Inches(0.5), Inches(6.92))
    n2 = nav_btn(s2, "Siguiente  →", Inches(11.5), Inches(6.92))
    appear_on_click(s2, card_shapes[0], 1)
    appear_on_click(s2, card_shapes[1], 2)
    appear_on_click(s2, card_shapes[2], 3)

    # ——— 3 calendario ———
    s3cal = prs.slides.add_slide(blank)
    set_bg(s3cal, VERDE_OSC)
    header(s3cal, "DIAGNÓSTICO  ·  OFERTA Y CARGA MES A MES")
    title(s3cal, "Cuándo actuar.", top=Inches(0.82))
    subc = s3cal.shapes.add_textbox(Inches(0.5), Inches(1.55), Inches(12.2), Inches(0.4))
    tf = subc.text_frame
    p = tf.paragraphs[0]
    p.clear()
    r = p.add_run()
    r.text = "Verde = pasto del mes.  Oro = lo que comen.  Rojo = ese mes hay que intervenir."
    r.font.size = Pt(16)
    r.font.color.rgb = CREMA
    r.font.name = "Calibri"

    cifras = [
        ("HOY", "6 vacas"),
        ("EL SUELO CUBRE EL AÑO", "5 en año seco"),
        ("DECISIÓN", "Bajar 1  o diferido"),
    ]
    cif_shapes = []
    for i, (lab, val) in enumerate(cifras):
        c = box(s3cal, Inches(0.5 + i * 4.2), Inches(2.05), Inches(4.0), Inches(1.15), CREMA, ORO)
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.16)
        tf.margin_top = Inches(0.1)
        p = tf.paragraphs[0]
        p.clear()
        r = p.add_run()
        r.text = lab
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = VERDE
        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        r2.text = val
        r2.font.size = Pt(20)
        r2.font.bold = True
        r2.font.color.rgb = VERDE_OSC
        cif_shapes.append(c)

    acc = [
        "Bajar carga o no subir",
        "Rotar · dividir el cuadro",
        "Acercar aguada (si > 400 m)",
        "Sembrar 1–2 ha de diferido",
    ]
    acc_shapes = []
    for i, line in enumerate(acc):
        c = box(s3cal, Inches(0.5 + (i % 2) * 6.35), Inches(3.45 + (i // 2) * 1.15), Inches(6.15), Inches(1.02), CREMA, ORO)
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.18)
        tf.margin_top = Inches(0.22)
        p = tf.paragraphs[0]
        p.clear()
        r = p.add_run()
        r.text = line
        r.font.size = Pt(18)
        r.font.bold = True
        r.font.color.rgb = VERDE
        acc_shapes.append(c)

    regl2 = s3cal.shapes.add_textbox(Inches(0.5), Inches(5.85), Inches(12.2), Inches(0.7))
    tf = regl2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.clear()
    r = p.add_run()
    r.text = "El cuello de botella es el invierno. Inversión mínima. OTBN I: no sembrar ni desmontar. El monte se usa, no se saca."
    r.font.size = Pt(15)
    r.font.color.rgb = ORO
    r.font.name = "Calibri"

    btn_backc = nav_btn(s3cal, "←  Atrás", Inches(0.5), Inches(6.92))
    n3 = nav_btn(s3cal, "Siguiente  →", Inches(11.5), Inches(6.92))
    for i, sh in enumerate(cif_shapes + acc_shapes):
        appear_on_click(s3cal, sh, i + 1)

    # ——— 4 ———
    s3 = prs.slides.add_slide(blank)
    set_bg(s3, VERDE_OSC)
    header(s3, "CÓMO SE USA  ·  UNA VISITA AL PRODUCTOR")
    title(s3, "Cuatro pasos.", top=Inches(0.78))

    pasos = [
        ("1", "Las 4 esquinas", "Plano de mensura: latitud Sur y longitud Oeste (φ y λ). No use X e Y."),
        ("2", "Las aguadas", "Toque cada represa o molino. 400 m cría · 800 m adulta. No se recorta bebida."),
        ("3", "Manejo", "Inicial, Mejorando e Ideal. El calendario dice hasta dónde subir y en qué mes actuar."),
        ("4", "El informe", "Puntos críticos. Corto y mediano plazo. En el celular se ve igual; no queda guardado."),
    ]
    step_shapes = []
    top = 1.72
    for n, tit, body in pasos:
        c = box(s3, Inches(0.5), Inches(top), Inches(8.55), Inches(1.05), CREMA, ORO)
        circ = s3.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.68), Inches(top + 0.24), Inches(0.52), Inches(0.52))
        fill(circ, VERDE)
        tf = circ.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.clear()
        r = p.add_run()
        r.text = n
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = BLANCO
        tb = s3.shapes.add_textbox(Inches(1.4), Inches(top + 0.08), Inches(7.45), Inches(0.9))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.clear()
        r = p.add_run()
        r.text = tit
        r.font.size = Pt(17)
        r.font.bold = True
        r.font.color.rgb = VERDE
        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        r2.text = body
        r2.font.size = Pt(12)
        r2.font.color.rgb = TIERRA
        step_shapes.append(c)
        top += 1.12

    side = box(s3, Inches(9.25), Inches(1.72), Inches(3.6), Inches(4.18), CREMA, ORO)
    tf = side.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_top = Inches(0.16)
    p = tf.paragraphs[0]
    p.clear()
    r = p.add_run()
    r.text = "Para recordar"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = VERDE
    for line in [
        "Capas: satélite, INTA, OTBN, lluvia, forraje.",
        "Manejo: Inicial / Mejorando / Ideal.",
        "Calendario: verde pasto, rojo actuar.",
        "IA: interfaz, no el núcleo.",
        "Computadora: se puede guardar.",
        "Celular: vista rápida, no guarda.",
    ]:
        p = tf.add_paragraph()
        p.space_before = Pt(8)
        r = p.add_run()
        r.text = "▸ " + line
        r.font.size = Pt(12)
        r.font.color.rgb = TIERRA

    link = nav_btn(s3, "Abrir el gemelo", Inches(9.25), Inches(6.05), w=Inches(3.6), h=Inches(0.48))
    link.click_action.hyperlink.address = URL
    btn_back3 = nav_btn(s3, "←  Atrás", Inches(0.5), Inches(6.92))

    n1.click_action.target_slide = s2
    n2.click_action.target_slide = s3cal
    n3.click_action.target_slide = s3
    btn_back2.click_action.target_slide = s1
    btn_backc.click_action.target_slide = s2
    btn_back3.click_action.target_slide = s3cal

    appear_on_click(s3, step_shapes[0], 1)
    appear_on_click(s3, step_shapes[1], 2)
    appear_on_click(s3, step_shapes[2], 3)
    appear_on_click(s3, step_shapes[3], 4)

    prs.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
